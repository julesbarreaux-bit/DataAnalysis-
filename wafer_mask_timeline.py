"""
8" wafer preparation timeline + mask cleaning timeline (matplotlib, static PNG).

Reads directly from "I-PEX U-LEAP V2.xlsx":
  - "Wafer preparation schedule" sheet -> 8" wafers (ID starting with "8")
    that have at least one date recorded.
  - "Cleaning Schedule" sheet -> mask cleaning events, keyed by the mask's
    full name as written in the sheet (not abbreviated).

Both panels share the same x-axis (date) so they can be compared directly.
"""

import re
from datetime import timedelta
from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import openpyxl

XLSX_PATH = Path(__file__).with_name("I-PEX U-LEAP V2.xlsx")
OUTPUT_PATH = Path(__file__).with_name("charts") / "wafer_8in_mask_cleaning_timeline.png"

# Fixed categorical palette (colorblind-validated), used in a fixed slot order
# everywhere so a given role always gets the same color.
BLUE, AQUA, YELLOW, GREEN, VIOLET, RED, MAGENTA, ORANGE = (
    "#2a78d6", "#1baf7a", "#eda100", "#008300",
    "#4a3aa7", "#e34948", "#e87ba4", "#eb6834",
)

WAFER_STAGE_COLUMNS = [
    ("Wafer Start Date", BLUE),
    ("PZT Deposition Date", AQUA),
    ("TE Date", YELLOW),
    ("PVI Date", GREEN),
    ("BG Date", VIOLET),
    ("LoP Date", RED),
    ("DRIE Date", MAGENTA),
    ("Target ETD (shipped Date I-Pex)", ORANGE),
]

MASK_FAMILY_COLORS = {
    "BE": BLUE, "CAV": AQUA, "LE": YELLOW, "LOP": GREEN,
    "PL": VIOLET, "PVI": RED, "PZT": MAGENTA, "TE": ORANGE,
}


def mask_family(mask_name: str) -> str:
    prefix = re.split(r"-", mask_name, maxsplit=1)[0]
    return re.sub(r"\d+$", "", prefix).strip().upper()


def normalize_mask_name(raw: str) -> str:
    return re.sub(r"\s*-\s*", "-", raw.strip())


def load_wafer_rows(wb):
    ws = wb["Wafer preparation schedule"]
    headers = [c.value for c in ws[1]]
    col_idx = {h: headers.index(h) + 1 for h, _ in WAFER_STAGE_COLUMNS if h in headers}
    id_col, status_col = 1, 2

    rows = []
    for r in range(3, ws.max_row + 1):
        wafer_id = ws.cell(row=r, column=id_col).value
        if not wafer_id or not str(wafer_id).strip().startswith("8"):
            continue
        stages = {h: ws.cell(row=r, column=c).value for h, c in col_idx.items()}
        if not any(stages.values()):
            continue
        rows.append({
            "id": str(wafer_id).strip(),
            "status": ws.cell(row=r, column=status_col).value,
            "stages": stages,
        })
    return rows


def load_cleaning_events(wb):
    ws = wb["Cleaning Schedule"]
    events = []
    for r in range(2, ws.max_row + 1):
        mask = ws.cell(row=r, column=1).value
        date = ws.cell(row=r, column=3).value
        cleaner = ws.cell(row=r, column=4).value
        if mask and date:
            events.append((normalize_mask_name(mask), date, cleaner))
    return events


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    wafers = load_wafer_rows(wb)
    cleaning_events = load_cleaning_events(wb)

    wafer_dates = [d for w in wafers for d in w["stages"].values() if d]
    domain_min = min(wafer_dates) - timedelta(days=7)
    domain_max = max(wafer_dates) + timedelta(days=7)

    cleaning_events = [e for e in cleaning_events if domain_min <= e[1] <= domain_max]

    masks = {}
    for mask, date, cleaner in cleaning_events:
        masks.setdefault(mask, []).append((date, cleaner))
    mask_names = sorted(masks, key=lambda m: (mask_family(m), m))

    fig, (ax_wafer, ax_mask) = plt.subplots(
        2, 1, figsize=(15, 4 + 0.32 * (len(wafers) + len(mask_names))),
        sharex=True, gridspec_kw={"height_ratios": [max(len(wafers), 2), max(len(mask_names), 2)]},
    )

    # ---- Panel 1: wafer preparation stages ----
    active_stage_cols = [(h, c) for h, c in WAFER_STAGE_COLUMNS
                          if any(w["stages"].get(h) for w in wafers)]

    for i, w in enumerate(wafers):
        ax_wafer.axhline(i, color="#e1e0d9", linewidth=0.8, zorder=0)
        present = [(h, w["stages"][h]) for h, _ in active_stage_cols if w["stages"].get(h)]
        present.sort(key=lambda hp: hp[1])
        for h, d in present:
            color = dict(active_stage_cols)[h]
            ax_wafer.scatter(d, i, color=color, s=90, zorder=2,
                              edgecolors="white", linewidths=1.2)

    ax_wafer.set_yticks(range(len(wafers)))
    ax_wafer.set_yticklabels([f'{w["id"]}  ({w["status"]})' for w in wafers], fontsize=9)
    ax_wafer.set_ylim(-1, len(wafers))
    ax_wafer.invert_yaxis()
    ax_wafer.set_title("8 inch wafer", fontsize=12, loc="left")
    ax_wafer.grid(axis="x", color="#e1e0d9", linewidth=0.8)
    for spine in ("top", "right", "left"):
        ax_wafer.spines[spine].set_visible(False)

    stage_handles = [
        plt.Line2D([0], [0], marker="o", linestyle="", markersize=8,
                   markerfacecolor=c, markeredgecolor="white", label=h)
        for h, c in active_stage_cols
    ]
    ax_wafer.legend(handles=stage_handles, loc="upper left", bbox_to_anchor=(1.01, 1),
                     fontsize=9, frameon=False)

    # ---- Panel 2: mask cleaning events (full mask names, not abbreviated) ----
    for i, mask in enumerate(mask_names):
        color = MASK_FAMILY_COLORS.get(mask_family(mask), "#898781")
        ax_mask.axhline(i, color="#e1e0d9", linewidth=0.8, zorder=0)
        dates = [d for d, _ in masks[mask]]
        ax_mask.scatter(dates, [i] * len(dates), color=color, s=36, zorder=2,
                         edgecolors="white", linewidths=0.8)

    ax_mask.set_yticks(range(len(mask_names)))
    ax_mask.set_yticklabels(mask_names, fontsize=8.5)
    ax_mask.set_ylim(-1, len(mask_names))
    ax_mask.invert_yaxis()
    ax_mask.set_title("Mask cleaning", fontsize=12, loc="left")
    ax_mask.grid(axis="x", color="#e1e0d9", linewidth=0.8)
    for spine in ("top", "right", "left"):
        ax_mask.spines[spine].set_visible(False)

    family_handles = [
        plt.Line2D([0], [0], marker="o", linestyle="", markersize=7,
                   markerfacecolor=c, markeredgecolor="white", label=fam)
        for fam, c in MASK_FAMILY_COLORS.items()
        if any(mask_family(m) == fam for m in mask_names)
    ]
    ax_mask.legend(handles=family_handles, loc="upper left", bbox_to_anchor=(1.01, 1),
                    fontsize=9, frameon=False, title="Mask family")

    ax_mask.xaxis.set_major_locator(mdates.MonthLocator())
    ax_mask.xaxis.set_major_formatter(mdates.DateFormatter("%b %y"))
    fig.autofmt_xdate()

    fig.suptitle('Wafer preparation schedule — I-PEX U-LEAP V2.xlsx', fontsize=13, fontweight="bold", x=0.01, ha="left")
    fig.tight_layout(rect=(0, 0, 0.86, 0.97))

    OUTPUT_PATH.parent.mkdir(exist_ok=True)
    fig.savefig(OUTPUT_PATH, dpi=150)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
